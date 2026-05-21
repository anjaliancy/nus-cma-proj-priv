import math
import sys
from pathlib import Path

import cvxpy as cp
import numpy as np
import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / 'src'))

from cma.output_summary import export_milp_output_summary
from cma.port import Port, PortGraph, PortPool
from cma.servicegraph import ServiceGraph
from cma.serviceline import ServiceLine
from cma.vessel import Vessel, VesselPool


def _mock_problem():
	port_kwargs = {
		'fit_vessel_ranks': {1: 10},
		'berth_productivity': {1: 1000.0},
		'cost_portcall': {1: 1000.0},
		'cost_transship': 50.0,
		'cost_storage': 10.0,
		'transshipment_capacity': True,
		'max_draft': 20.0,
		'max_daily_call': 5,
	}
	p1 = Port('P1', 'Port 1', 120.0, 1.0, **port_kwargs)
	p2 = Port('P2', 'Port 2', 121.0, 2.0, **port_kwargs)
	portpool = PortPool([p1, p2])
	portgraph = PortGraph(
		ports_pool=portpool,
		mat_distance=np.array([[0.0, 2400.0], [2400.0, 0.0]]),
		mat_demand=np.array([[0.0, 100.0], [0.0, 0.0]]),
		filter_by_demand=False,
	)
	vessel = Vessel(
		v_rank=1,
		v_class=(100, 200),
		capacity=1000.0,
		draft=10.0,
		daily_chartering_cost=10000.0,
		bunkering_cost_coefs=[
			{'speed': float(speed), 'consumption': 20.0 + speed}
			for speed in np.arange(10.0, 19.0, 1.0)
		],
		unit_bunkering_cost=500.0,
	)
	vesselpool = VesselPool([vessel], [10])
	line = ServiceLine('TestLine', [p1, p2])
	line.set_buffer_profile([12.0, 12.0], [14.0, 14.0], False)
	od_pairs = [(0, 1)]
	od_pair_paths = [[line.get_shortest_path(p1, p2)]]
	return ServiceGraph([line]), [line], portgraph, vesselpool, od_pairs, od_pair_paths


def _shared_segment_problem():
	port_kwargs = {
		'fit_vessel_ranks': {1: 10, 2: 10},
		'berth_productivity': {1: 1000.0, 2: 1000.0},
		'cost_portcall': {1: 1000.0, 2: 1000.0},
		'cost_transship': 50.0,
		'cost_storage': 10.0,
		'transshipment_capacity': True,
		'max_draft': 20.0,
		'max_daily_call': 5,
	}
	p1 = Port('P1', 'Port 1', 120.0, 1.0, **port_kwargs)
	p2 = Port('P2', 'Port 2', 121.0, 2.0, **port_kwargs)
	portpool = PortPool([p1, p2])
	portgraph = PortGraph(
		ports_pool=portpool,
		mat_distance=np.array([[0.0, 240.0], [240.0, 0.0]]),
		mat_demand=np.array([[0.0, 900.0], [0.0, 0.0]]),
		filter_by_demand=False,
	)
	vessel_small = Vessel(
		v_rank=1,
		v_class=(100, 500),
		capacity=500.0,
		draft=10.0,
		daily_chartering_cost=1000.0,
		bunkering_cost_coefs=[
			{'speed': float(speed), 'consumption': 10.0 + speed}
			for speed in np.arange(10.0, 19.0, 1.0)
		],
		unit_bunkering_cost=100.0,
	)
	vessel_large = Vessel(
		v_rank=2,
		v_class=(500, 1000),
		capacity=1000.0,
		draft=10.0,
		daily_chartering_cost=2000.0,
		bunkering_cost_coefs=[
			{'speed': float(speed), 'consumption': 12.0 + speed}
			for speed in np.arange(10.0, 19.0, 1.0)
		],
		unit_bunkering_cost=100.0,
	)
	vesselpool = VesselPool([vessel_small, vessel_large], [10, 10])
	line_used = ServiceLine('UsedLine', [p1, p2])
	line_unused = ServiceLine('UnusedLine', [p1, p2])
	for line in [line_used, line_unused]:
		line.set_buffer_profile([12.0, 12.0], [14.0, 14.0], False)
	od_pairs = [(0, 1)]
	od_pair_paths = [[line_used.get_shortest_path(p1, p2)]]
	return ServiceGraph([line_used, line_unused]), portgraph, vesselpool, od_pairs, od_pair_paths


def _tuneparams():
	return {
		'turnon-transship_shipclass_restriction': 0,
		'turnon-vessel_speed_optimization': 0,
		'turnon-port_operations_constraint': 1,
		'turnon-transit_time_penalty': 0,
		'turnon-schedule_adherence': 0,
		'ctrparam-kts_buffer': 0,
		'ctrparam-transship_A': 100,
		'ctrparam-speed_soft_cap_kts': 16.5,
		'ctrparam-speed_penalty_multiplier': 2.0,
		'ctrparam-transit_penalty_multiplier': 1000.0,
		'ctrparam-buffer_penalty_below_15pct': 1000.0,
		'ctrparam-buffer_penalty_above_30pct': 2000.0,
		'unfulfilled_demand_penalty': 1e6,
		'BigM-transship': 10000,
		'BigM-n_ships': 10,
		'BigM-saildays': 100,
		'BigM-line_capacity': 100000,
		'BigM-portcall_cost': 1e9,
		'solver-verbose': False,
	}


def _skip_if_no_usable_mip_solver():
	for solver in [cp.SCIP, cp.GLPK_MI, cp.CBC, cp.GUROBI]:
		if solver not in cp.installed_solvers():
			continue
		x = cp.Variable(boolean=True)
		prob = cp.Problem(cp.Minimize(x), [x >= 0])
		try:
			prob.solve(solver=solver, verbose=False)
		except Exception:
			continue
		if x.value is not None:
			return
	pytest.skip('No usable MILP solver is available for solver-backed diagnostics test')


def _fake_solution():
	return {
		'total cost': 10000.0,
		'chartering cost': 1000.0,
		'transshipment cost': 200.0,
		'bunkering cost': 3000.0,
		'portcall cost': 400.0,
		'transit penalty cost': 5000.0,
		'unfulfilled demand penalty cost': 0.0,
		'buffer penalty cost': 400.0,
		'kpi_teus_input': 100.0,
		'kpi_teus_fulfilled': 100.0,
		'kpi_teus_delayed': 0.0,
		'kpi_avg_delay_days': 0.0,
		'kpi_lines_buffer_above_30': 1,
		'kpi_lines_buffer_below_15': 0,
		'solver_status': 'mock',
		'solver_name': 'mock',
		'solver_solve_time': 0.0,
		'solver_mip_gap': 0.0,
		'solver_best_bound': 10000.0,
		'solver_obj_val': 10000.0,
		'solver_node_count': 0,
		'line diagnostics': [{
			'line_index': 0,
			'line_name': 'TestLine',
			'selected_week': 2.0,
			'vessel_count': 2.0,
			'ship_allocation': [2.0],
			'selected_speed_kts': 14.0,
			'weekly_capacity_teu': 1000.0,
			'line_capacity_total_teu': 2000.0,
			'segment_flows': [
				{'from_port': 'P1', 'to_port': 'P2', 'flow_teu': 100.0},
				{'from_port': 'P2', 'to_port': 'P1', 'flow_teu': 0.0},
			],
			'port_stay_days': [0.125, 0.125],
			'line_sailing_days': 13.75,
			'buffer_value': 0.2,
			'buffer_violation_lb_hours': 0.0,
			'buffer_violation_ub_hours': 0.2,
			'chartering_cost': 1000.0,
			'transshipment_cost': 200.0,
			'bunkering_cost': 3000.0,
			'portcall_cost': 400.0,
			'buffer_penalty_lb': 0.0,
			'buffer_penalty_ub': 400.0,
			'buffer_penalty_cost': 400.0,
		}],
	}


def test_fulfill_demands_returns_line_diagnostics_and_reconciles_costs():
	_skip_if_no_usable_mip_solver()
	servicegraph, _, portgraph, vesselpool, od_pairs, od_pair_paths = _mock_problem()
	solution = servicegraph.fulfill_demands(
		od_pairs,
		od_pair_paths,
		portgraph,
		vesselpool,
		week_levels=[1, 2, 3, 4],
		tuneparams=_tuneparams(),
	)

	diagnostics = solution['line diagnostics']
	assert len(diagnostics) == 1
	line_diag = diagnostics[0]
	for key in [
		'selected_week',
		'vessel_count',
		'ship_allocation',
		'selected_speed_kts',
		'weekly_capacity_teu',
		'segment_flows',
		'port_stay_days',
		'line_sailing_days',
		'buffer_value',
		'chartering_cost',
		'transshipment_cost',
		'bunkering_cost',
		'portcall_cost',
		'buffer_penalty_cost',
	]:
		assert key in line_diag

	assert math.isclose(
		sum(item['chartering_cost'] for item in diagnostics),
		solution['chartering cost'],
		rel_tol=1e-6,
		abs_tol=1e-6,
	)
	assert math.isclose(
		sum(item['transshipment_cost'] for item in diagnostics),
		solution['transshipment cost'],
		rel_tol=1e-6,
		abs_tol=1e-6,
	)
	assert math.isclose(
		sum(item['bunkering_cost'] for item in diagnostics),
		solution['bunkering cost'],
		rel_tol=1e-6,
		abs_tol=1e-6,
	)
	assert math.isclose(
		sum(item['portcall_cost'] for item in diagnostics),
		solution['portcall cost'],
		rel_tol=1e-6,
		abs_tol=1e-6,
	)
	assert math.isclose(
		line_diag['segment_flows'][0]['flow_teu'],
		100.0,
		rel_tol=1e-6,
		abs_tol=1e-5,
	)
	assert 'capacity_proxy_flow_teu' in line_diag['segment_flows'][0]

	reconciled = (
		solution['chartering cost']
		+ solution['transshipment cost']
		+ solution['bunkering cost']
		+ solution['portcall cost']
		+ solution['transit penalty cost']
		+ solution['unfulfilled demand penalty cost']
		+ solution['buffer penalty cost']
	)
	assert math.isclose(solution['total cost'], reconciled, rel_tol=1e-6, abs_tol=1e-5)


def test_shared_physical_segment_capacity_is_not_pooled_between_services():
	_skip_if_no_usable_mip_solver()
	servicegraph, portgraph, vesselpool, od_pairs, od_pair_paths = _shared_segment_problem()
	tuneparams = _tuneparams()
	tuneparams.update({
		'turnon-vessel_speed_optimization': 1,
		'unfulfilled_demand_penalty': 1e9,
	})
	solution = servicegraph.fulfill_demands(
		od_pairs,
		od_pair_paths,
		portgraph,
		vesselpool,
		week_levels=[1],
		tuneparams=tuneparams,
	)

	used_line_diag = solution['line diagnostics'][0]
	used_leg_flow = used_line_diag['segment_flows'][0]['flow_teu']
	used_capacity = used_line_diag['weekly_capacity_teu']
	assert math.isclose(used_leg_flow, 900.0, rel_tol=1e-6, abs_tol=1e-5)
	assert used_capacity + 1e-5 >= used_leg_flow


def test_export_milp_output_summary_workbook(tmp_path):
	_, service_lines, portgraph, vesselpool, _, _ = _mock_problem()
	solution = _fake_solution()
	output_path = tmp_path / 'milp_output_summary.xlsx'

	export_milp_output_summary(
		solution,
		service_lines,
		portgraph,
		vesselpool,
		proforma_metadata={'TestLine': {'service_type': 'OWN'}},
		output_path=output_path,
	)

	assert output_path.exists()
	wb = openpyxl.load_workbook(output_path, data_only=True)
	assert set(wb.sheetnames) == {'summary', 'run_metadata'}
	summary = wb['summary']
	headers = [cell.value for cell in summary[1]]
	assert 'linename' in headers
	assert 'buffer_penalty' in headers
	line_names = [summary.cell(row=row, column=1).value for row in range(2, summary.max_row + 1)]
	assert line_names == ['TestLine', '__TOTAL__']

	idx_bunker = headers.index('bunkeringcost') + 1
	idx_charter = headers.index('charteringcost') + 1
	idx_portcall = headers.index('portcallcost') + 1
	idx_transship = headers.index('transshipmentcost') + 1
	assert summary.cell(row=3, column=idx_bunker).value == summary.cell(row=2, column=idx_bunker).value
	assert summary.cell(row=3, column=idx_charter).value == summary.cell(row=2, column=idx_charter).value
	assert summary.cell(row=3, column=idx_portcall).value == summary.cell(row=2, column=idx_portcall).value
	assert summary.cell(row=3, column=idx_transship).value == summary.cell(row=2, column=idx_transship).value
