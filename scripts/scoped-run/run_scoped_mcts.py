"""
Run MCTS (not just a baseline solve) on the same scoped subset used by
run_scoped_baseline.py -- every VSA line, every FIX line, plus the specific
OWN lines named in the original client feedback (BBX2CNC, BBX3CNC, BMXCNC,
CP2CNC, YCXCNC) -- then export the best network found through the standard
output_summary workbook (summary / run_metadata / cargo_flow_routes /
changes sheets), same format used for the full-network runs.

Scoped rather than the full 31-line/182-port network because that full
network is known to take 2+ hrs just to build the optimisation problem
(see docs/progress_summary_2026-09-11.md) -- this subset solves the baseline
in ~17s, so MCTS epochs are actually tractable here.

Resumable search (prof feedback - tree persistence): by default this script
saves its search tree to --tree-file after every run and reloads it on the
next run, via MonteCarloTree.commit_one_step()/save_tree()/load_tree() in
cma.mcts. Each run commits at most one step - the first action on the path
to the best node found so far - and forgets the sibling branches that lost
out at that decision point, so a later run can't re-explore them. Pass
--fresh to ignore any saved tree and start over from the baseline network.

Usage (from repo root, using the project venv):
    .venv\\Scripts\\python.exe scripts\\scoped-run\\run_scoped_mcts.py
    .venv\\Scripts\\python.exe scripts\\scoped-run\\run_scoped_mcts.py --epochs 15 --depth 3
    .venv\\Scripts\\python.exe scripts\\scoped-run\\run_scoped_mcts.py --fresh
"""
import sys, os, time, argparse

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, "src"))

ap = argparse.ArgumentParser()
ap.add_argument("--out", default=None, help="output .xlsx path")
ap.add_argument("--extra-lines", nargs="*", default=[
    "BBX2CNC", "BBX3CNC", "BMXCNC", "CP2CNC", "YCXCNC",
])
ap.add_argument("--epochs", type=int, default=15)
ap.add_argument("--depth", type=int, default=3)
ap.add_argument("--weight", type=float, default=0.74)
ap.add_argument("--timelimit", type=int, default=120)
ap.add_argument("--mipgap", type=float, default=0.10)
ap.add_argument("--seed", type=int, default=7)
ap.add_argument("--tree-file", default=None,
    help="where the search tree is saved/resumed from (default: "
         "data/output/scoped_mcts_tree.pkl)")
ap.add_argument("--fresh", action="store_true",
    help="ignore any saved tree at --tree-file and start a brand new search "
         "from the baseline network, discarding all previously committed history")
args = ap.parse_args()

import random
import numpy as np
random.seed(args.seed); np.random.seed(args.seed)

from cma.data_reader import (read_vessel_class_data, read_port_data, read_sailing_distance_data,
    read_demand_with_transit_time, read_cnc_proforma_data)
from cma.port import PortGraph
from cma.servicegraph import ServiceGraph
from cma.mcts import MonteCarloTree
from cma.output_summary import (export_milp_output_summary, build_modified_lines_and_cost_delta,
    write_change_summary_sheet)

def log(msg=""):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)

tree_file = args.tree_file or os.path.join(ROOT, "data", "output", "scoped_mcts_tree.pkl")

log("Loading data...")
vesselpool = read_vessel_class_data()
portpool_main, _ = read_port_data()
dist_matrix = read_sailing_distance_data(portpool_main)
demand_matrix, transit_time_matrix = read_demand_with_transit_time(portpool_main)
proforma = read_cnc_proforma_data(portpool_main, vesselpool, dist_matrix=dist_matrix)
all_lines = proforma["lines"]
proforma_metadata = proforma.get("metadata")

wanted_names = set(args.extra_lines)
scoped_lines = [
    line for line in all_lines
    if line.service_type in ("VSA", "FIX") or line.name() in wanted_names
]
scoped_line_names = sorted(l.name() for l in scoped_lines)
log(f"Scoped network: {len(scoped_lines)} of {len(all_lines)} lines "
    f"({sum(1 for l in scoped_lines if l.service_type=='VSA')} VSA, "
    f"{sum(1 for l in scoped_lines if l.service_type=='FIX')} FIX, "
    f"{sum(1 for l in scoped_lines if l.service_type=='OWN')} OWN)")
log("Lines: " + ", ".join(scoped_line_names))

portgraph = PortGraph(portpool_main, dist_matrix, demand_matrix,
                      mat_transit_time=transit_time_matrix, filter_by_demand=False)
servicegraph = ServiceGraph(scoped_lines)

WEEK_LEVELS = [1, 2, 3, 4, 5, 6, 7]
TP = {'turnon-transship_shipclass_restriction': 0, 'turnon-vessel_speed_optimization': 1,
 'turnon-tight_line_capacity_linearization': 1, 'turnon-port_operations_constraint': 1,
 'turnon-transit_time_penalty': 1, 'turnon-demand_fulfillment_cap': 1, 'ctrparam-kts_buffer': 0,
 'ctrparam-transship_A': 100, 'ctrparam-speed_soft_cap_kts': 16.5, 'ctrparam-speed_penalty_multiplier': 2.0,
 'ctrparam-transit_penalty_multiplier': 1000.0, 'ctrparam-buffer_penalty_below_15pct': 1000.0,
 'ctrparam-buffer_penalty_above_30pct': 5000.0, 'unfulfilled_demand_penalty': 1e6, 'BigM-transship': 10000,
 'BigM-n_ships': 10, 'BigM-saildays': 100, 'BigM-line_capacity': 30000, 'BigM-portcall_cost': 1e7,
 'turnon-schedule_adherence': 0, 'schedule_buffer_hrs': 120.0, 'solver-MIPGap': args.mipgap,
 'solver-TimeLimit': args.timelimit, 'solver-MIPFocus': 1, 'solver-verbose': False}

# Always solve the true, never-modified scoped baseline - this is the fixed
# reference the final "changes" sheet diffs against, regardless of whether
# this run resumes a saved tree that's already several commits deep.
log("Solving baseline (root) ...")
t0 = time.time()
baseline_solution = servicegraph.solve_approximated(
    portgraph, vesselpool, min_cost=True, week_levels=WEEK_LEVELS, tuneparams_2=TP,
)
true_baseline_cost = servicegraph.total_cost()
baseline_service_lines = servicegraph.tolist_serviceLine()
log(f"Baseline root cost = {true_baseline_cost:,.2f}  (solved in {time.time()-t0:.0f}s)")

tree = None
resumed_from_saved_tree = False
if not args.fresh and os.path.exists(tree_file):
    log(f"Found saved tree at {tree_file}, attempting to resume ...")
    loaded = MonteCarloTree.load_tree(tree_file)
    loaded_names = sorted(l.name() for l in loaded.root_node.graph.tolist_serviceLine())
    if loaded_names != scoped_line_names:
        log("Saved tree's line set doesn't match this run's --extra-lines - "
            "can't safely resume, starting fresh instead.")
    else:
        tree = loaded
        resumed_from_saved_tree = True
        prior_commits = tree.committed_action_trace()
        log(f"Resumed: {len(prior_commits)} step(s) already committed in earlier runs, "
            f"current root cost = {tree.root_node.graph.total_cost():,.0f}")
        for i, text in enumerate(prior_commits, 1):
            log(f"  committed {i}. {' '.join(text.split())}")

if tree is None:
    if args.fresh:
        log("Starting fresh search from the baseline network (--fresh).")
    else:
        log("No saved tree found - starting fresh search from the baseline network.")
    tree = MonteCarloTree(servicegraph=servicegraph, portgraph=portgraph, vesselpool=vesselpool,
        discount_fac=0.5, valid_weight_proportion=args.weight, max_depth=args.depth, c_param=1e-2,
        min_cost=True, week_levels=WEEK_LEVELS, milp_tuneparams=TP)

root_cost = tree.root_node.graph.total_cost()

log("-" * 70)
t_start = time.time()
for ep in range(1, args.epochs + 1):
    te = time.time()
    try:
        tree.run(1, display=False)
        best = tree.get_best_node()
        rc = tree.root_node.graph.total_cost()
        bc = best.graph.total_cost()
        impr = rc - bc
        pct = (impr / rc * 100) if rc else 0.0
        cum = (time.time() - t_start) / 60
        log(f"epoch {ep:>3}/{args.epochs} | {time.time()-te:>5.0f}s (cum {cum:>5.1f}m) | "
            f"root={rc:,.0f} best={bc:,.0f} impr={impr:,.0f} ({pct:.2f}%) "
            f"depth={best.get_depth()} nodes={tree.total_number_of_nodes()}")
    except Exception as e:
        log(f"epoch {ep:>3}/{args.epochs} | ERROR: {e!r}")

log("-" * 70)
committed_this_run = tree.commit_one_step()
if committed_this_run is not None:
    log(f"Committed one step this run: {' '.join(tree.committed_action_trace()[-1].split())}")
    log(f"New root cost = {tree.root_node.graph.total_cost():,.0f} "
        f"(tree re-rooted here; sibling branches discarded)")
else:
    log("No improving action found yet this run - nothing committed, "
        "tree state saved as-is for the next run.")

tree.save_tree(tree_file)
log(f"Search tree saved to: {tree_file}")

best = tree.get_best_node()
best_cost = best.graph.total_cost()
log(f"FINAL: current root={tree.root_node.graph.total_cost():,.0f}  best={best_cost:,.0f}  "
    f"true baseline={true_baseline_cost:,.0f}  "
    f"total improvement vs. true baseline={true_baseline_cost-best_cost:,.0f} "
    f"({(true_baseline_cost-best_cost)/true_baseline_cost*100:.2f}%)")
log("Remaining best-node action trace (beyond the current committed root):")
action_trace_lines = []
for i, act in enumerate(best.trace_actions(), 1):
    try:
        text = act.explain(portgraph, tree.root_node.graph, i).replace("\n", " ")
    except Exception:
        text = str(act)
    log(f"  {i}. {text}")
    action_trace_lines.append(text)

log("Re-solving best network to recover full MILP diagnostics ...")
solution = best.graph.solve_approximated(
    portgraph, vesselpool, min_cost=True, week_levels=WEEK_LEVELS, tuneparams_2=TP,
)
service_lines = best.graph.tolist_serviceLine()

out_path = args.out or os.path.join(ROOT, "data", "output",
    f"scoped_mcts_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
written = export_milp_output_summary(
    solution, service_lines, portgraph, vesselpool,
    proforma_metadata=proforma_metadata, output_path=out_path,
)

import pandas as pd
# NOTE: pd.ExcelWriter(written, engine='openpyxl') in default ('w') mode
# replaces the WHOLE workbook, silently dropping any sheet not re-written
# here (found the hard way: this pattern, copied from
# scripts/export_mcts_best_network.py, was dropping 'cargo_flow_routes').
# Read every existing sheet first and re-write all of them together.
cargo_flow_routes_df = pd.read_excel(written, sheet_name='cargo_flow_routes')
metadata_df = pd.read_excel(written, sheet_name='run_metadata')
all_committed = tree.committed_action_trace()
extra_rows = pd.DataFrame([
    {'metric': 'scenario_note', 'value':
        'Scoped MCTS best-network summary (VSA+FIX+named OWN lines, not the full 31-line network).'},
    {'metric': 'mcts_true_baseline_cost', 'value': true_baseline_cost},
    {'metric': 'mcts_root_cost', 'value': root_cost},
    {'metric': 'mcts_best_cost', 'value': best_cost},
    {'metric': 'mcts_improvement_pct', 'value': round((true_baseline_cost - best_cost) / true_baseline_cost * 100, 2)},
    {'metric': 'mcts_epochs', 'value': args.epochs},
    {'metric': 'mcts_depth', 'value': args.depth},
    {'metric': 'mcts_action_trace', 'value': ' | '.join(action_trace_lines) or '(no changes selected)'},
    {'metric': 'mcts_tree_file', 'value': tree_file},
    {'metric': 'mcts_resumed_from_saved_tree', 'value': resumed_from_saved_tree},
    {'metric': 'mcts_committed_steps_total', 'value': len(all_committed)},
    {'metric': 'mcts_committed_action_trace', 'value': ' | '.join(
        ' '.join(text.split()) for text in all_committed) or '(none committed yet)'},
])
metadata_df = metadata_df[metadata_df['metric'] != 'scenario_note']
metadata_df = pd.concat([metadata_df, extra_rows], ignore_index=True)
summary_df = pd.read_excel(written, sheet_name='summary')
with pd.ExcelWriter(written, engine='openpyxl') as writer:
    summary_df.to_excel(writer, sheet_name='summary', index=False)
    metadata_df.to_excel(writer, sheet_name='run_metadata', index=False)
    cargo_flow_routes_df.to_excel(writer, sheet_name='cargo_flow_routes', index=False)

modified_lines, cost_rows = build_modified_lines_and_cost_delta(
    baseline_solution, baseline_service_lines, solution, service_lines,
)
import openpyxl
wb = openpyxl.load_workbook(written)
if 'changes' in wb.sheetnames:
    del wb['changes']
write_change_summary_sheet(wb.create_sheet('changes'), modified_lines, cost_rows)
wb.save(written)

log(f"Workbook written to: {written}")
